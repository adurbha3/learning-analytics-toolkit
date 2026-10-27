"""Builds a multi-tab Excel reporting workbook: a student risk summary with
a pandas pivot-table-style category breakdown and conditional formatting,
and a marketing overview with a conversion-rate chart and a color-scaled
keyword opportunity table."""
from __future__ import annotations

import pandas as pd
from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

HEADER_FILL = "4472C4"


def _write_dataframe(ws: Worksheet, df: pd.DataFrame, start_row: int = 1, start_col: int = 1) -> tuple[int, int]:
    """Writes a DataFrame as a formatted table; returns (end_row, end_col)."""
    for j, column_name in enumerate(df.columns):
        cell = ws.cell(row=start_row, column=start_col + j, value=column_name)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor=HEADER_FILL)
    for i, row in enumerate(df.itertuples(index=False), start=1):
        for j, value in enumerate(row):
            ws.cell(row=start_row + i, column=start_col + j, value=value)
    return start_row + len(df), start_col + len(df.columns) - 1


def _add_color_scale(ws: Worksheet, column_index: int, start_row: int, end_row: int, low: str, high: str) -> None:
    column_letter = get_column_letter(column_index)
    rule = ColorScaleRule(start_type="min", start_color=low, end_type="max", end_color=high)
    ws.conditional_formatting.add(f"{column_letter}{start_row}:{column_letter}{end_row}", rule)


def _autofit_columns(ws: Worksheet) -> None:
    for column_cells in ws.columns:
        values = [cell.value for cell in column_cells if cell.value is not None]
        length = max((len(str(v)) for v in values), default=0)
        ws.column_dimensions[column_cells[0].column_letter].width = min(40, max(10, length + 2))


def build_workbook(
    student_scores: pd.DataFrame,
    conversion_by_source: pd.DataFrame,
    keyword_opportunities: pd.DataFrame,
    output_path: str,
) -> str:
    wb = Workbook()

    ws1 = wb.active
    ws1.title = "Student Risk Summary"
    end_row, end_col = _write_dataframe(ws1, student_scores)
    risk_col = student_scores.columns.get_loc("risk_score") + 1
    _add_color_scale(ws1, risk_col, 2, end_row, low="63BE7B", high="F8696B")

    category_summary = (
        student_scores.pivot_table(index="category", values="risk_score", aggfunc=["count", "mean"])
    )
    category_summary.columns = ["student_count", "avg_risk_score"]
    category_summary = category_summary.round(1).reset_index()
    summary_col = end_col + 3
    summary_end_row, _ = _write_dataframe(ws1, category_summary, start_row=1, start_col=summary_col)

    risk_chart = BarChart()
    risk_chart.title = "Risk Score by Student"
    risk_chart.y_axis.title = "Risk Score"
    data = Reference(ws1, min_col=risk_col, min_row=1, max_row=end_row)
    categories = Reference(
        ws1, min_col=student_scores.columns.get_loc("student_id") + 1, min_row=2, max_row=end_row
    )
    risk_chart.add_data(data, titles_from_data=True)
    risk_chart.set_categories(categories)
    ws1.add_chart(risk_chart, f"{get_column_letter(summary_col)}{summary_end_row + 3}")

    ws2 = wb.create_sheet("Marketing Overview")
    conv_end_row, conv_end_col = _write_dataframe(ws2, conversion_by_source)

    conversion_chart = BarChart()
    conversion_chart.title = "Conversion Rate by Source"
    conversion_chart.y_axis.title = "Conversion Rate"
    conv_data = Reference(
        ws2, min_col=conversion_by_source.columns.get_loc("conversion_rate") + 1, min_row=1, max_row=conv_end_row
    )
    conv_categories = Reference(
        ws2, min_col=conversion_by_source.columns.get_loc("source_medium") + 1, min_row=2, max_row=conv_end_row
    )
    conversion_chart.add_data(conv_data, titles_from_data=True)
    conversion_chart.set_categories(conv_categories)
    ws2.add_chart(conversion_chart, f"A{conv_end_row + 3}")

    keyword_col = conv_end_col + 4
    top_keywords = keyword_opportunities.head(10)
    keyword_end_row, _ = _write_dataframe(ws2, top_keywords, start_row=1, start_col=keyword_col)
    opportunity_col = keyword_col + top_keywords.columns.get_loc("opportunity_score")
    _add_color_scale(ws2, opportunity_col, 2, keyword_end_row, low="FFFFFF", high="63BE7B")

    for ws in (ws1, ws2):
        _autofit_columns(ws)

    wb.save(output_path)
    return output_path
