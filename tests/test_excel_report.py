import pandas as pd
from openpyxl import load_workbook

from app.excel_report import build_workbook


def _sample_frames():
    student_scores = pd.DataFrame([
        {"student_id": "S001", "avg_pre_score": 84.8, "score_trend": 3.3,
         "avg_engagement": 5.0, "risk_score": 11.2, "category": "on_track"},
        {"student_id": "S002", "avg_pre_score": 44.0, "score_trend": -4.0,
         "avg_engagement": 1.0, "risk_score": 73.4, "category": "needs_support"},
    ])
    conversion_by_source = pd.DataFrame([
        {"source_medium": "google / organic", "total_sessions": 200, "total_conversions": 20,
         "avg_bounce_rate": 0.3, "conversion_rate": 0.1},
        {"source_medium": "facebook / cpc", "total_sessions": 100, "total_conversions": 2,
         "avg_bounce_rate": 0.6, "conversion_rate": 0.02},
    ])
    keyword_opportunities = pd.DataFrame([
        {"keyword": "sat prep online", "search_volume": 5000, "current_ranking": None,
         "click_through_rate": None, "campaign": "unranked", "is_ranked": False, "opportunity_score": 5000.0},
        {"keyword": "math tutor near me", "search_volume": 3000, "current_ranking": 45,
         "click_through_rate": 0.02, "campaign": "organic", "is_ranked": True, "opportunity_score": 750.0},
    ])
    return student_scores, conversion_by_source, keyword_opportunities


def test_build_workbook_creates_expected_sheets(tmp_path):
    student_scores, conversion, keywords = _sample_frames()
    output_path = build_workbook(student_scores, conversion, keywords, str(tmp_path / "report.xlsx"))

    wb = load_workbook(output_path)
    assert wb.sheetnames == ["Student Risk Summary", "Marketing Overview"]


def test_build_workbook_writes_student_rows_correctly(tmp_path):
    student_scores, conversion, keywords = _sample_frames()
    output_path = build_workbook(student_scores, conversion, keywords, str(tmp_path / "report.xlsx"))

    wb = load_workbook(output_path)
    ws = wb["Student Risk Summary"]
    header = [cell.value for cell in ws[1]][:6]
    assert header == list(student_scores.columns)
    assert ws.cell(row=2, column=1).value == "S001"
    assert ws.cell(row=3, column=1).value == "S002"


def test_build_workbook_adds_charts_to_both_sheets(tmp_path):
    student_scores, conversion, keywords = _sample_frames()
    output_path = build_workbook(student_scores, conversion, keywords, str(tmp_path / "report.xlsx"))

    wb = load_workbook(output_path)
    assert len(wb["Student Risk Summary"]._charts) == 1
    assert len(wb["Marketing Overview"]._charts) == 1


def test_build_workbook_adds_conditional_formatting(tmp_path):
    student_scores, conversion, keywords = _sample_frames()
    output_path = build_workbook(student_scores, conversion, keywords, str(tmp_path / "report.xlsx"))

    wb = load_workbook(output_path)
    assert len(list(wb["Student Risk Summary"].conditional_formatting)) >= 1
    assert len(list(wb["Marketing Overview"].conditional_formatting)) >= 1
